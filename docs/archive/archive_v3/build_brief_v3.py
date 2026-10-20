from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "Silent_Avalon_Project_Brief.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.70), Inches(.70)
sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)
sec.footer_distance = Inches(.30)

def font_style(style, size, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + a), "Arial")
    fonts.set(qn("w:eastAsia"), "Heiti SC")
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + a), None)
    for el in list(style.element.iter(qn("w:pBdr"))):
        el.getparent().remove(el)

for name, size, bold in [
    ("Normal", 10.5, False), ("Title", 22, True), ("Subtitle", 12, False),
    ("Heading 1", 14, True), ("Heading 2", 11.5, True),
    ("List Bullet", 10.5, False), ("List Number", 10.5, False),
]:
    font_style(doc.styles[name], size, bold)

normal = doc.styles["Normal"].paragraph_format
normal.line_spacing = 1.18
normal.space_after = Pt(6)
normal.widow_control = True
for name in ("Heading 1", "Heading 2"):
    fmt = doc.styles[name].paragraph_format
    fmt.space_before = Pt(13 if name == "Heading 1" else 8)
    fmt.space_after = Pt(6)
    fmt.keep_with_next = True
doc.styles["Title"].paragraph_format.space_after = Pt(6)
doc.styles["Subtitle"].paragraph_format.space_after = Pt(12)

def p(text, boldlead=None, style=None):
    para = doc.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        para.add_run(boldlead).bold = True
        para.add_run(text[len(boldlead):])
    else:
        para.add_run(text)
    return para

def h(text, level=1):
    return doc.add_heading(text, level)

def table(headers, rows, widths, keep=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        borders.append(el)
    pr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, val in (("top", "85"), ("bottom", "85"), ("left", "105"), ("right", "105")):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:w"), val)
        e.set(qn("w:type"), "dxa")
        margins.append(e)
    pr.append(margins)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_idx, values in enumerate([headers] + rows):
        row = t.rows[0] if row_idx == 0 else t.add_row()
        cant = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant)
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.13
            if keep and row_idx < len(rows):
                para.paragraph_format.keep_with_next = True
            run = para.add_run(value)
            run.font.size = Pt(9.5)
            if row_idx == 0:
                run.bold = True
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E6EDF3")
                cell._tc.get_or_add_tcPr().append(shade)
            if j == 0 and len(value) < 16:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.line_spacing = 1
    gap.add_run().font.size = Pt(2)
    return t

def ref(n, text, url):
    para = p(f"[{n}] {text}")
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.1
    for r in para.runs:
        r.font.size = Pt(9)
    para.add_run(" ")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "24567A")
    rpr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    r.append(rpr)
    tx = OxmlElement("w:t")
    tx.text = "原文链接"
    r.append(tx)
    link.append(r)
    para._p.append(link)

doc.core_properties.title = "六人阿瓦隆中显式交流的作用"
doc.core_properties.subject = "Structured Communication in Avalon: Research Project Specification"
doc.core_properties.author = ""
doc.core_properties.keywords = "Avalon, MARL, Reinforcement Learning, Communication, CS7329"
doc.core_properties.comments = ""
p("六人阿瓦隆中显式交流的作用", style="Title")
p("Structured Communication in Avalon: Research Project Specification", style="Subtitle")
p("CS7329 Reinforcement Learning　｜　项目规划 v3（沉默基线＋结构化通信）　｜　2026年10月20日")

h("1 项目概述")
p("本项目研究六人隐藏身份游戏《阿瓦隆》中，在沉默对局的基础上加入结构化显式交流，会如何改变阵营胜率、策略行为及信息传递。固定角色为两名平民、莫甘娜、刺客、梅林和派西维尔。研究采用同一套游戏规则与强化学习算法：先训练不发言的沉默智能体作为基线，再在此基础上逐级开放通信槽位，训练可通信智能体；并在训练后关闭通信信道，测量可通信策略对信道的依赖程度。")
p("为什么以沉默为基线：沉默对局的决策更少，粗略估计每局约 60 个决策、完整通信约 160 个，待 pilot 实测；也不存在“消息被忽略”的 cheap-talk 学习难题，因此更容易训练成稳定、可复现的策略。它同时承担环境和 PPO 实现的正确性检查，可作为通信配置的初始化起点；如果通信学习失败，它本身也能构成完整的保底结果。研究问题因此表述为“交流带来了什么”，而不是“禁言夺走了什么”。两种表述下主胜率对比的数值相同、符号相反，但实验顺序、消融方向和结论的指向不同。", boldlead="为什么以沉默为基线：")
p("建议的核心方案是小型 recurrent PPO 多智能体训练，加上五槽结构化消息协议：身份声明、阵营判断、判断依据、组队建议、行动倾向。每个槽位都有独立开关，因此“交流”可以从沉默基线出发逐项加入，而不只是整体开或关。自然语言交流作为独立扩展实验，并复用同一套五槽结构。核心实验保留“显式交流对比沉默”的因果问题，能够利用本地 RTX 5070 Ti Laptop 完成开发和试验，并把学校 HPC 用于多随机种子训练。9B 模型不承担核心实验的端到端强化学习训练。", boldlead="建议的核心方案")
p("范围界定：五槽消息是对线下一次约 2 分钟、约 400 tokens 发言的结构化抽象。它覆盖这类发言的主要信息类型，但不包含措辞、语气、即兴追问和长篇推理，不能等同于开放式语言讨论。核心结论只适用于所定义的槽位、取值范围、角色配置和训练所得策略。若最终论文要直接声称“自然语言交流的作用”，必须补充自然语言实验，并单独报告其结果。", boldlead="范围界定：")
h("研究问题", 2)
table(
    ["问题", "需要回答的内容", "主要证据"],
    [
        ["RQ1\nGame Balance", "相对沉默基线，加入交流后哪一阵营获益更多？", "好人最终胜率、三次任务成功概率、条件刺杀成功率"],
        ["RQ2\nStrategic Behavior", "相对沉默基线，组队、投票、破坏与身份隐藏如何调整？", "按角色和局面分层的行为指标及训练轨迹"],
        ["RQ3\nImplicit Signaling", "沉默基线中，公开行为是否承担信息传递？加入显式交流后，这部分隐式信号被补充还是被挤出？", "信息增益、接收者干预、跨种子和跨搭档验证"],
        ["RQ4\nMessage Components", "从沉默基线逐项加入的五类信息中，哪些贡献了交流的效果？", "逐项加入的训练消融、测试时关闭槽位、槽位使用率与真实性、听者因果影响"],
    ], [1.12, 2.75, 2.84]
)
p("不预设交流一定有利于好人。交流可能帮助好人协调，也可能帮助坏人欺骗或识别梅林；任务阶段与刺杀阶段的效果还可能相互抵消。不同槽位的作用方向也可能相反，例如身份声明可能帮助派西维尔，也可能把梅林暴露给刺客。RQ1、RQ2 为核心研究目标；RQ4 至少完成“最小通信”与“完整通信”两档对照及测试时槽位消融；RQ3 为有明确证据标准的探索目标。即使没有显著胜率变化，或没有发现稳定信号，也可以形成有效结果。")

h("2 文献定位与课程关联")
p("DeepRole 研究五人阿瓦隆，明确不使用自然语言交流，以推断、搜索和学习的价值函数处理合作与对抗。[1] 因而“阿瓦隆可以在沉默条件下由智能体完成”已有先例，这也支持以沉默对局作为基线；本项目研究的是在同一训练框架下加入结构化通信带来的差异，而不是把沉默阿瓦隆本身作为首次提出的环境。")
p("Is it Percival time yet 是 SIGBOVIK 2017 的初步人类游戏分析，报告 66 局采集、剔除 5 局后的 61 局数据。[2] 它可帮助提出刺杀、身份暴露和破坏协调的指标，但样本与角色配置存在限制，不足以直接推导交流的因果效应。")
p("STRATEGIST 使用 LLM 生成和改进策略，并结合低层搜索进行决策，不能直接视为 PPO 式参数训练。[3] 尤其需要注意，附录 O 已描述六人局有对话与无对话的人类实验，但其角色为 Merlin、三名 Servants、Minion 与 Assassin，不含 Percival 与 Morgana，并且是 LLM 智能体与人类对局，不是受控训练。因此本项目的拟贡献是可复现的训练对照、策略适应分析及信号机制检验；单纯“六人局有无对话对照”不足以确立创新性。该文正文与附录的人类实验总局数表述不同，本文不把其汇总胜率当作可直接比较的基准。")
p("AvalonBench 及其代码库可作为环境实现的候选起点。[4,10] Lan 等人的六人 Avalon 研究采用包含 Percival 与 Morgana 的角色配置，并分析了协调、伪装和信息分享等行为，可用于借鉴行为分类。[5] 即使角色构成相同，具体规则、语言模型和实验规模仍须逐项核对，不应直接与本项目胜率拼成排行榜。")
p("结构化通信的先例与风险。把消息当作策略的离散动作、与其他动作一起用强化学习训练，可追溯到 Foerster 等人的 RIAL；本项目沿用这一思路，只是用 PPO 代替 DQN，并把一条消息拆成五个槽位。[16] RLupus 在狼人杀中用 PPO 训练村民的离散通信信道（狼人为规则策略），九人局村民胜率从无通信的 0.044 最高提到 0.45，而且信号长度与取值范围对性能的影响是非线性的。[11] CICERO 先由规划模块产生 intent，再由受控对话模型把 intent 转成语言。[12] 本项目的五槽消息与这种 intent 思路一致：RL 决定“说什么”，语言层只负责“怎么说”。Cao 等人发现，自利智能体难以利用没有内在含义的 cheap-talk 信道，合作型智能体则可以。[13] 阿瓦隆阵营内部合作、阵营之间对抗，而且坏人可以模仿任何消息，所以通信是否被学会需要单独诊断，不能只看槽位名称。", boldlead="结构化通信的先例与风险。")
p("拟贡献：以沉默对局为基线，提供一个在规则、算法和预算上受控、逐级加入通信的实验；用可独立开关的五槽协议把“交流”分解为可逐项加入的信息类型；把整体胜负分解为任务与刺杀两条路径；通过行为干预区分“动作包含信息”与“动作被用来影响其他智能体”。这是一项实验研究与机制分析项目，不以提出新的 RL 优化算法为必要条件。", boldlead="拟贡献：")
h("与课程内容的连接", 2)
p("课程中的 MDP 在这里扩展为 Partially Observable Stochastic Game。完整局面包括隐藏身份，单个玩家只能访问自己的 observation history。Actor 学习策略 π，critic 学习该玩家信息条件下的价值 V；终局胜负提供 return，TD 与 GAE 用于估计 advantage，PPO 用于更新策略。[7,8] Self-play 则使其他玩家的策略也随训练变化。")
p("课程项目要求动机与 RL 关联、相关研究、技术方法、文献或自有实验结果，以及局限性讨论。[C1] 本方案覆盖这些内容，并与仓库中的部分可观察性、函数近似、Actor–Critic 和 PPO 教学材料衔接。[C2]")

h("3 环境规则与信息边界")
h("核心六人配置", 2)
p("主实验固定为两名平民（Loyal Servants）、一名莫甘娜（Morgana）、一名刺客（Assassin）、一名梅林（Merlin）和一名派西维尔（Percival）。好人阵营为梅林、派西维尔及两名平民，坏人阵营为莫甘娜与刺客。沉默基线、各通信配置及训练后关闭信道的评估均采用这一配置；派西维尔与莫甘娜的身份辨别及混淆机制属于核心研究范围。")
table(
    ["项目", "固定规范"],
    [
        ["任务与表决", "五轮任务人数为 2、3、4、3、4；至少四票赞成才通过，平票否决。同一任务连续五次提案被否决，坏人获胜。"],
        ["任务执行", "好人只能提交 Success；坏人可选 Success 或 Fail。六人局每轮一张 Fail 即失败。公开失败票总数，不公开个人任务票。"],
        ["终局判定", "三次任务失败即坏人胜；三次成功后由 Assassin 刺杀，命中 Merlin 则坏人胜，否则好人胜。"],
        ["身份信息", "所有人知道自己的角色。Merlin 知道两名坏人的座位，但不知道 Percival 的位置；Percival 看到 Merlin 与 Morgana 的两名候选者，但不知道谁是谁；Morgana 与 Assassin 互认坏人队友；两名平民没有额外身份信息。"],
        ["随机化", "每局随机分配角色和初始队长，随后队长按规则轮换；禁用 Lady of the Lake 等附加机制。"],
    ], [1.10, 5.61]
)
p("上述规则以游戏规则书为基准。[9] 座位编号不能固定对应角色；Percival 的两名候选者使用无身份标签的集合表示，不能用排列顺序泄露谁是 Merlin。投票与任务牌必须先分别提交、再一起揭示，防止执行顺序泄露未公开选择。仿真器完整状态仅用于裁判和离线分析。")
h("结构化通信协议", 2)
p("发言时机：每次队长公布提案后、全员投票前有一轮公开讨论。从队长开始按座位顺序，六名玩家各发一条结构化消息，后发言者可以看到先发言者的内容；队长先发言，相当于说明组队理由。三次任务成功后进入刺杀阶段，两名坏人各有一次公开发言，只开放阵营判断槽，含义改为“我认为 j 是梅林”。每次提案只有一轮发言，比 v1 的提案前、提案后两轮少一半通信决策，以缩短每名玩家的决策序列。", boldlead="发言时机：")
p("每条消息最多包含五个槽位，对应线下一次约 2 分钟发言的五类内容。所有取值都是离散的，均由策略选择；允许虚假内容，裁判只检查合法性，不检查真假。")
table(
    ["槽位", "线下对应", "离散取值", "约束"],
    [
        ["① 身份声明", "我是谁、我自称什么角色", "不声明、忠臣、派西维尔、梅林", "允许虚假声明；可在不同轮次改口，历次声明公开可见"],
        ["② 阵营判断", "我认为谁好谁坏", "对其余五名玩家各选：不评、好、坏", "允许虚假判断；刺杀阶段改为标记梅林候选"],
        ["③ 判断依据", "我为什么这样判断", "对每名被评为好或坏的玩家选一条依据：任务记录、投票记录、组队记录（各附所指轮次）、身份视野声明，或无依据", "依附于 ②；公开记录类依据只能指向真实发生的事件，但可以选择性引用；身份视野声明不可验证"],
        ["④ 组队建议", "我点哪些人上车", "不建议，或当前任务人数下的任一合法队伍（最多 20 种）", "使用与队长提案动作相同的队伍编码；队长此槽固定为其提案"],
        ["⑤ 行动倾向", "我赞成或反对，希望谁回应", "投票倾向：不表态、赞成、反对；点名：无，或本轮尚未发言的玩家", "倾向不约束实际投票；被点名者发言时可见点名记录"],
    ], [1.0, 1.25, 2.3, 2.16]
)
p("开关与“不说”：每个槽位有独立开关。开启的槽位都包含“不说”选项，所以五句话中的任何一句都可以不说；关闭的槽位固定为“不说”，不进入策略的动作损失。③ 依附于 ②，开启 ③ 时必须开启 ②；①、④、⑤ 与其他槽位相互独立。", boldlead="开关与“不说”：")
p("判断依据的约束：公开记录类依据只能引用确有其事的事件，由合法动作掩码保证。理由有二：线下歪曲公开记录会被当场识破；把可验证的公开依据与不可验证的身份视野声明分开，便于检验“有依据的指控”是否更能影响听者。若 pilot 发现该约束使 ③ 几乎不被使用，可把“允许虚假引用”作为事先说明的变体，单独报告。", boldlead="判断依据的约束：")
p("渲染与记录：每条消息以槽位元组形式存储，并用固定模板渲染成中文或英文句子，供日志阅读和自然语言扩展使用。例如：“我是派西维尔。我认为 3 号是坏人，因为第 2 轮任务失败时他在队里。我建议 1、2、5 号上车。我反对当前队伍，请 4 号回应。”完整五槽消息渲染后约为数十到一百余 tokens，低于 400 tokens 的线下发言上限；差额对应措辞、语气和冗余，不属于核心实验。", boldlead="渲染与记录：")
p("沉默基线 S 关闭全部五个槽位并跳过讨论阶段，对智能体而言等价于每名玩家每个槽位都选择“不说”。公开的队伍、已揭示投票、任务结果和失败票数仍然可见。训练后关闭信道（F → S）同样实现为强制所有槽位“不说”，而不是引入训练中从未出现的新 NULL 符号，以减轻分布外输入效应。所有条件都禁止私聊、跨玩家共享记忆及利用延迟、报错等额外信道。")
h("通信配置与消融", 2)
table(
    ["配置", "身份", "判断", "依据", "组队", "倾向", "用途"],
    [
        ["S", "关", "关", "关", "关", "关", "沉默基线（对照组）"],
        ["J（可选）", "关", "开", "关", "关", "关", "只加入阵营判断"],
        ["M", "关", "开", "开", "关", "关", "最小通信：加入谁好谁坏、为什么"],
        ["M+Claim", "开", "开", "开", "关", "关", "逐项加入：身份声明"],
        ["M+Team", "关", "开", "开", "开", "关", "逐项加入：组队建议"],
        ["M+Intent", "关", "开", "开", "关", "开", "逐项加入：行动倾向"],
        ["F", "开", "开", "开", "开", "开", "完整结构化通信"],
    ], [0.95, 0.6, 0.6, 0.6, 0.6, 0.6, 2.76], keep=True
)
p("配置按从沉默到完整通信的阶梯排列。核心层为 S、M、F 三个训练配置；扩展层为在 M 上逐项加入一个槽位的三个配置，衡量每类信息在最小通信之上的增量。与之互补，测试时在 F 上关闭单个槽位，衡量完整通信对该槽位的依赖，无须重新训练；两种消融的结论不一致，说明槽位之间存在交互。可选的 J 位于 S 与 M 之间：② 与 ③ 在 M 及以上的配置中总是同时出现，若不训练 J，就无法单独识别“判断依据”在训练适应后的作用；至少应在 M 的冻结 checkpoint 上做一次关闭 ③ 的测试时消融。各配置使用相同的网络结构、输出头、发言顺序和训练预算，只通过掩码改变可用槽位。")
p("这种通信安排是有带宽限制的实验协议。记录槽位定义、取值范围、发言顺序和机会数并保持不变；不能把结果直接推广到任意长度、任意时间的自由讨论。给槽位和取值起了自然语言名称，也不代表智能体学到的含义一定与名称相同；含义需要用第 6 节的真实性与听者影响指标实测。")

h("4 强化学习方法")
h("策略与价值网络", 2)
p("采用 recurrent Independent PPO 作为首选基线：好人与坏人各使用一套小型 MLP＋GRU actor–critic，同阵营共享参数，输入包含自身角色和其合法私有信息。每个玩家独立维护 recurrent hidden state；共享网络权重不意味着共享其他玩家的记忆。先使用 128 维编码与 128 维 GRU 作为试验配置，并在正式比较前冻结架构。")
p("输入为公共游戏历史、自身历史行动、角色允许的信息及通信记录；通信记录按槽位编码，并区分“槽位关闭”与“选择不说”。输出使用分阶段 action heads：队长选择合法队伍，全员选择赞成或反对，坏人选择任务牌，Assassin 选择刺杀目标。采用 legal action mask 排除非法动作。六人局合法队伍最多 20 种，可直接枚举，无须生成任意文本动作。")
p("消息输出头：讨论阶段为五个槽位分别设输出头。① 是一个 4 类头；② 是五个按玩家分解的 3 类头；③ 对每名被评为好或坏的玩家设一个依据类型头和一个轮次头，由掩码保证只能选真实事件；④ 与队长提案共享队伍编码；⑤ 是投票倾向头加点名头。一条消息的对数概率等于已开启子动作对数概率之和，PPO 比率按这一联合概率计算。所有配置保留全部输出头，关闭的槽位只通过掩码移除，以保证各配置网络容量一致。消息头与行动头共享参数和编码器，属于 RIAL 式训练；不采用让梯度穿过信道的 DIAL，因为阿瓦隆的听者中有对手。[16]", boldlead="消息输出头：")
p("选择 recurrent policy 是为了利用历史信息，而非把当前观察错误地当成完整 Markov state。核心 critic 同样只接收该玩家可获得的信息；若日后采用 centralized critic，必须单独说明训练时的特权信息，且保证执行 actor 看不到它。")
h("奖励与更新", 2)
p("每个玩家在己方最终获胜时得到 +1、失败时得到 −1，中间奖励为 0；以有限回合未折扣终局收益为目标，取 γ＝1，避免新增通信步骤改变同一胜负的折扣权重。两阵营分别汇总损失，避免四名好人相对两名坏人造成隐含的训练权重差异。")
p("学习循环为：按当前策略采集完整对局 → 按玩家整理观察、合法动作、奖励和记忆序列 → 用价值估计与 GAE 计算 advantage → 用 PPO 更新行动头、消息头及 critic → 保存 checkpoint 并重复。通信动作也通过终局奖励学习。关闭槽位、强制“不说”和其他强制动作不计入 actor 的可选动作损失；padding 不参与梯度。")
p("核心实验不增加“诚实发言”“成功传信”或“隐瞒身份”的额外奖励，避免把希望发现的策略直接写入奖励。是否出现信息交换和隐藏行为，由胜负激励下的学习结果决定。PPO、GAE 的具体超参数先在独立 pilot 中确定，之后各通信配置共用；不能按测试胜率分别挑选最有利参数。")
h("训练顺序：沉默基线优先", 2)
p("先单独训练沉默基线 S，直到其对随机与规则对手的胜率稳定，再开始任何通信配置。这一步同时检验环境、奖励、掩码与 PPO 实现；S 不过关时不训练通信配置。S 与通信配置使用同一套在 pilot 中确定的超参数和相同的总训练局数；不能因为 S 容易训练而减少其预算，否则基线偏弱，会高估交流的作用。")
p("分支训练：每个随机种子先训练 S 至总预算的一半，得到共享的分支点 checkpoint。S 从该点继续训练剩余一半；M、F 等通信配置从同一点加入消息头后，继续训练相同局数。新增的消息输入权重初始化为零，使分支时的通信智能体在行为上与 S 完全相同；消息输出头从均匀分布起步。这样各配置的总训练局数相同，同一种子下的差异只来自分支后能否交流，可以做配对比较。分支后，各配置的历史对手池只收录本配置的 checkpoint。", boldlead="分支训练：")
p("热启动的代价是：通信智能体可能沿用沉默阶段已形成的隐式约定，从而低估从头学习时交流的作用。因此至少保留一个从头训练的 F 种子作为对照，并在论文中报告两者的差异。若 pilot 中分支设计导致消息头完全不被使用，则改为所有配置从头训练、预算相同。")
h("通信学习诊断", 2)
p("结构化本身不能保证通信被学会。任何 cheap-talk 博弈都存在“说话者随意发言、听者完全忽略”的均衡，独立学习者容易停在这里。[13] 结构化的作用是缩小组合空间、让组队建议与提案共享编码、让依据锚定在真实事件上，并使分析可以逐槽进行。为此设置三项检查：")
p("玩具检查：把坏人固定为规则策略，只训练好人在 M 配置下通信。若梅林的阵营判断仍不能提高好人胜率，优先排查实现错误，而不是得出“交流无用”的结论。", boldlead="玩具检查：", style="List Number")
p("发送端诊断：计算各槽取值与发言者真实身份、被评价者真实阵营之间的互信息，按角色分层。", boldlead="发送端诊断：", style="List Number")
p("接收端诊断：在冻结 checkpoint 上替换某条消息中的某个槽位，测量听者投票与组队分布的 KL 变化，即 causal influence。[6,14]", boldlead="接收端诊断：", style="List Number")
p("第二阶段结束时，若 M 与 F 的接收端影响接近零，按第 9 节的检查点处理，不直接进入正式训练。可选的初始化办法是：先用“诚实规则发言者”的行为克隆初始化消息头，再用 RL 微调，使槽位名称在训练初期与其含义一致；之后是否转向欺骗由胜负奖励决定。若采用，所有通信配置统一使用，并在论文中报告。")
h("Self play 与非平稳性", 2)
p("先训练两个阵营的当前策略，再维护历史 checkpoint 池，减少只适应最新对手的循环。建议 pilot 尝试每局以 50% 概率使用当前对手、25% 使用历史对手、25% 使用固定规则对手；正式实验锁定混合比例。对手的固定轨迹不参与当前策略的 on-policy 更新，同局队友仍保持独立观察。")
p("保留随机合法策略和固定规则策略作 sanity baselines。规则策略只利用合法信息，例如根据任务失败记录累计嫌疑、坏人以固定概率破坏、刺客按简单公开行为评分猜 Merlin。它们用于检查学习是否改善，不能代表人类或最优策略。IPPO 在此环境的有效性需要试验确认，不能由其在合作基准上的表现直接推断。[7]")

h("5 实验矩阵与统计计划")
table(
    ["实验", "训练配置", "评估配置", "解释"],
    [
        ["S → S", "沉默", "沉默", "沉默基线（对照组）"],
        ["M → M", "最小通信", "最小通信", "加入阵营判断与判断依据"],
        ["F → F", "完整通信", "完整通信", "主处理组：完整结构化通信"],
        ["M+k → M+k", "逐项加入配置", "同一配置", "在 M 上加入槽位 k 的增量（扩展层）"],
        ["F → S\nM → S", "F 或 M", "冻结权重后全部槽位强制“不说”", "通信依赖：训练好的通信策略失去信道时的表现"],
        ["F → F−k", "完整通信", "冻结权重后关闭槽位 k", "测试时槽位依赖，k 取身份、依据、组队、倾向"],
        ["S → S mask", "沉默", "削弱指定历史线索", "RQ3：基线中隐式信号的接收者依赖"],
    ], [1.0, 1.35, 1.75, 2.61]
)
p("主比较为 F → F 相对 S → S，衡量在沉默基线上加入完整结构化交流的作用；M → M 相对 S → S 衡量“谁好谁坏、为什么”这一最小通信的作用；F → F 相对 M → M 衡量身份声明、组队建议和行动倾向的合计增量。它们构成从沉默到完整通信的阶梯。")
p("F → S 与 F → F−k 使用与 F → F 完全相同的 checkpoint 和评估初始条件，只改变可用槽位。它们度量通信依赖，可能包含分布外输入效应，不等价于沉默基线，也不等价于重新学会在缺少该槽位时游戏。若 F → S 明显低于 S → S，说明通信策略依赖信道，失去信道后也不再具备沉默基线的打法。反方向的 S → F 不做：沉默策略从未见过消息，在评估时打开信道没有意义。")
p("训练预算：先完成每个配置 2,000 局的功能与吞吐 pilot；依据耗时和学习曲线，在每次独立训练 20,000–100,000 局的候选范围内选择固定预算 B。该范围是规划参数，不是已验证的收敛需求。核心层 S、M、F 每个配置 3 个独立随机种子，共 9 次训练，资源允许再增至每配置 5 个；扩展层三个逐项加入配置各 3 个种子，再加 9 次训练；可选 J 配置与一个从头训练的 F 种子最后考虑。采用第 4 节的分支训练时，同一种子的各配置共享前半段 S 训练，实际计算量少于独立训练。种子少于 3 个的配置只作描述性报告。")
p("各配置匹配对局数、网络容量、PPO 更新次数与超参数，报告有效决策数、消息动作数和实际训练耗时。开启的槽位越多，消息子动作越多，各配置的信息和动作预算本来就不同；匹配对局数不等于完全匹配计算量。因此同时给出性能对训练局数及实际时间的曲线。")
p("评估预算：每个训练种子在其对应的全部评估条件下各冻结评估 1,000 局。核心层每组种子约 10 个评估条件，即 S → S；M → M、M → S、M 关闭 ③；F → F、F → S 及四个 F → F−k，三个种子合计约 3 万局。配对使用相同角色分配和初始队长表，策略采样使用独立随机流；没有训练更新或跨局记忆。训练种子是主要重复单位，不能把评估局数当成同样数量的独立训练模型。")
p("统计报告：先报告每个种子的指标，再汇总配对差值、效应量和 95% 区间。主结果使用按训练种子配对的分层 bootstrap，局内重复事件按完整游戏归组；种子数较少时明确区间不稳定。固定模型的单项胜率可另给 Wilson 区间，但不能用它替代训练不确定性。预先把好人最终胜率设为主指标，其余为机制与探索指标。")
p("同训练群体 self-play 只能说明该群体下的经验平衡。增加不同训练种子的好人和坏人交叉对战，以及与固定对手池的评估，以区分策略强弱、共同适应和通信效果。对阵双方都使用被测环境规则，且分别汇报每一阵营的表现。")

h("6 指标与可解释的结果")
h("RQ1 游戏平衡", 2)
p("核心指标为 ΔGood＝P(Good wins | F → F) − P(Good wins | S → S)，即相对沉默基线加入交流后好人胜率的变化。正值表示在该实验设定下交流更有利于好人，负值表示更有利于坏人。同样报告 M → M 相对 S → S 的差，以及 F → F 相对 M → M 的差。50% 只是阵营胜率对称的参考点，不意味着这款非对称游戏必须天然公平。")
p("为解释总胜率，分解为 P(Good wins)＝P(达到三次任务成功) × [1 − P(刺杀命中 Merlin | 达到三次任务成功)]。同时报告三次任务失败和连续五次否决导致的坏人获胜比例。刺杀成功率的分母只能是实际进入刺杀阶段的对局。")
p("任务成功率按任务序号报告，分母为实际执行该轮任务的次数，并报告到达该轮的局数；不能把未发生的后续任务记作失败。刺杀条件样本很少时报告宽区间，必要时按预先规定对所有条件等量增加评估。已知两名坏人位置时，在四名好人中均匀随机猜 Merlin 的命中率是 25%，可作刺杀参考基线。")
h("RQ2 策略行为", 2)
table(
    ["行为", "建议量化方式", "解释限制"],
    [
        ["组队与投票", "提议及通过队伍的好人比例；不同角色的赞成率；对含坏人队伍的赞成率", "按轮次、队伍人数、否决次数分层；最后机会的赞成不一定是判断错误"],
        ["任务破坏", "坏人在队时的 Fail 率；双坏人在队时零张、一张、两张 Fail 的分布", "放过任务可能用于伪装，两次破坏也不能自动判为协调失败"],
        ["梅林隐藏", "仅用公开历史训练身份 probe，评估识别 Merlin 的 log loss；结合任务与终局收益", "probe 难以识别不等于实际刺客一定被欺骗"],
        ["辨别与混淆", "Percival 对两名候选者的组队和投票偏好；Morgana 对这些偏好的影响", "控制局面与候选者历史；被支持不能直接等同于被认作 Merlin"],
        ["刺客推断", "条件刺杀准确率；对历史投票或组队记录的消融敏感度", "控制进入刺杀阶段的样本选择差异"],
    ], [1.0, 3.30, 2.41]
)
p("角色真值只在离线分析中用于计算这些指标。行为比较须覆盖早期、中期、最终 checkpoint；选取案例时使用明确抽样规则，同时展示支持与不支持假设的对局，不只挑选有趣的成功案例。")
h("RQ4 消息组件", 2)
table(
    ["分析", "建议量化方式", "解释限制"],
    [
        ["槽位使用", "各槽选择“不说”的比例随训练的变化，按角色、阵营和轮次分层", "使用率高不等于有用；可能只是无代价的习惯"],
        ["真实性", "身份声明与真实角色的一致率；被评为“坏”的玩家确为坏人的比例；梅林是否隐瞒、坏人是否冒充", "槽位含义由行为实测，不以名称为准"],
        ["依据作用", "各依据类型的使用分布；公开记录依据、身份视野声明与无依据三种判断对听者投票的影响", "② 与 ③ 同时出现，需配合测试时关闭 ③ 或 J 配置"],
        ["听者影响", "替换单个槽位后听者投票与组队分布的 KL 变化", "影响大不等于对己方有利，须结合收益"],
        ["胜率贡献", "逐项加入（M+k 相对 M）与测试时关闭（F−k 相对 F）的胜率差，按阵营分别报告", "前者是槽位在最小通信上的增量，后者是完整通信对它的依赖；两者不一致说明槽位间有交互"],
    ], [1.0, 3.30, 2.41]
)
p("另检查显式通道是否挤出隐式信号：以沉默基线 S 中投票序列对隐藏身份的额外预测能力为参照，比较 M、F 下的同一指标。如果开启显式通道后投票变得更不含信息，说明两种信道存在替代关系。")

h("7 隐式信号的证据标准")
p("本节的分析以沉默基线 S 为主：没有显式信道时，公开行为是唯一可能的信息传递途径，S 中测得的隐式信号水平也是 RQ3 后半问“被补充还是被挤出”的参照。公开行为与隐藏角色相关，只能证明行为包含可推断信息。要声称出现通信，还应证明接收者利用了这些行为；要进一步声称存在策略性信号，则需要发送者行为对接收者和收益的系统性作用。已有 emergent communication 研究说明，单看相关性或任务成绩容易产生误判。[6]")
p("第一层 信息性：使用独立测试对局训练和评估简单 probe。比较“公共局面＋队伍与任务结果”与“再加入个人投票序列”的隐藏身份预测效果。数据按完整对局和训练种子划分，避免同一对局片段进入训练与测试。预测提升只记为额外信息，不直接标为隐式通信。", boldlead="第一层 信息性：")
p("第二层 接收者使用：在固定 checkpoint 上回放同一接收者历史，将其他玩家的历史投票身份遮蔽，但保留总票数、提案结果与当前合法动作；从修改后的历史重新计算 GRU 状态，再比较接收者行动分布和后续收益。使用同长度占位作为 placebo，并以训练时见过遮蔽输入的辅助条件检验分布外影响。仅在推理时遮蔽导致性能下降，仍不足以单独证明通信。", boldlead="第二层 接收者使用：")
p("第三层 策略性信号：优先分析不改变表决通过结果的非关键投票。在可合法实现、保持票数与通过结果一致的反事实投票组合之间比较接收者反应，并检查信息持有者是否系统性选择能影响队友的模式。进一步改变搭档或其响应策略，观察发送者策略是否随之调整。若只完成前两层，应使用“行为线索被利用”的结论，避免宣称已经证明发送者意图或形成完整通信协议。", boldlead="第三层 策略性信号：")
p("重点分析 Percival 如何利用两名候选者的公开行为，以及 Morgana 是否形成稳定的混淆模式。稳健性检查包括角色与座位随机置换、未见过的角色分配和跨种子 cross-play。信号在不同种子间不兼容不等于群体内部没有信号。任务牌身份不公开，只能分析公开失败票数及其上下文。")
p("最低完成标准：至少完成信息性 probe 和一个可解释的接收者干预，清楚报告证据层级。若第三层证据不足，把 RQ3 作为探索性观察与后续研究方向，仍保留完整的 RQ1、RQ2 贡献。")

h("8 算力安排与自然语言扩展")
h("本地与 HPC 分工", 2)
p("本地 RTX 5070 Ti Laptop 优先运行规则引擎、并行环境、小型 recurrent PPO 和统计分析；先计时测量 CPU 仿真与 GPU 更新各自占比，不能默认瓶颈一定在 GPU。正式开跑前记录实际显存、RAM、功耗限制和软件版本，本规划不假设已验证该设备的吞吐或训练时长。")
p("建议为核心训练设置 40–60 GPU 小时的初始总预算上限，并用 pilot 修订 B：预计训练时间约为 配置数 × 种子数 × B × 每局平均训练耗时，另加调参与评估；核心层为 3 × 3，扩展层再加 3 × 3。若超预算，先降低单次训练量并保留重复种子；HPC 用于独立种子任务与扩展评估。核心项目不依赖持续取得 HPC 配额。")
p("AvalonBench 代码库的 README 说明目前只支持五人局（Merlin、两名 Servants、Minion、Assassin），不含 Percival 与 Morgana，页面也未标明许可证。[10] 因此首选按已固定规则实现轻量状态机，AvalonBench 只作为规则与接口参考。课程仓库提供算法学习与实现参考，但不是现成 Avalon 环境。此阶段不要求完整复现 DeepRole 或 STRATEGIST。")
h("语言层与自然语言扩展", 2)
p("分工：沿用 CICERO 把“说什么”和“怎么说”分开的思路。[12] RL 策略输出五槽结构化消息；语言层只做两件事：把结构化消息渲染成自然语言（模板即可，LLM 可选），以及在自然语言条件下把他人的发言解析回五槽结构。语言层不需要懂阿瓦隆策略，冻结的 API 模型或本地小模型即可胜任，无须专门训练。核心实验完全不依赖 LLM。", boldlead="分工：")
p("自然语言扩展采用一个冻结的本地 9B 模型，或一个固定版本的低成本 API 模型，分别扮演六个独立玩家。每名玩家使用隔离的上下文和记忆；只共享模型权重。发言以 JSON 输出五个槽位，关闭的槽位从输出格式中删除，另可附不超过 400 tokens 的自由文本，对应线下约 2 分钟的发言。NL-F、NL-M、NL-S 分别对应 F、M、S 的槽位开关，其他决策提示与预算保持一致。该实验验证语言代理中的方向是否类似，不能作为 MARL 策略已经学会自然语言的证据。")
p("先用 20 局 pilot 检查规则理解、格式合法率、上下文长度、单局 token 和时延，再视预算每条件完成 100–300 局。所有游戏动作通过结构化接口选择；对非法输出使用固定重试与回退规则并报告频率。不能因模型表现不佳而丢弃整局。")
p("本地 RTX 5070 Ti Laptop 为 12 GB 显存。9B 参数的原始权重在 FP16 下约需 18 GB，4-bit 下约需 4.5 GB，尚未包含量化元数据、KV cache 和运行开销。4-bit 量化推理是否可行须实测；能推理不代表同一设备可承担六角色长轨迹的端到端 LLM RL 训练。优先单模型串行或受限批处理，避免常驻六份权重。")
p("API 成本按实际总输入与输出 token 分别乘以执行时价格计算。400 tokens 的发言上限会让上下文随对局增长，单局 token 须在 pilot 中实测。初始扩展预算可设为不超过 50 美元，超限则减少局数或转本地推理。本项目不需要先确定商业模型或付费方案。自由语言结果应单独展示，不能与结构化通信实验混合平均。")
p("关于蒸馏开源模型：蒸馏不作为本课程交付，列为后续工作。蒸馏的前提是教师比学生更会玩，而现有通用 LLM 的阿瓦隆水平有限：AvalonBench 报告 ChatGPT 扮演好人、对阵规则 bot 坏人时胜率为 22.2%，规则 bot 好人为 38.2%。[4] 直接蒸馏 API 模型的发言会复制其弱点，而且多数商业 API 的服务条款限制用其输出训练其他模型，使用前须核对。更合理的两条路线：一是以本项目训练出的 RL 策略为教师做 policy distillation，[15] 把结构化观察和模板化发言转成文本，监督微调开源模型输出同样的五槽决策；二是用许可宽松的开源模型自博弈，只保留获胜方轨迹做筛选式行为克隆并迭代，这本身是一种基于结果的策略改进。12 GB 显存下，1–4B 模型的 LoRA 或 QLoRA 监督微调可行，9B 模型的 RL 微调不现实。", boldlead="关于蒸馏开源模型：")
p("如果研究重点必须从第一天就是“自然语言＋可训练 RL”，可另选冻结 LLM 编码器加小型 RL 决策头的混合路线，但需额外设计消息理解、缓存与训练接口。这会改变时间和算力预算，建议作为后续项目，而非本课程最低交付条件。")

h("9 进度交付与失败处理")
table(
    ["阶段", "建议时间", "验收内容"],
    [
        ["规则与试运行", "10月20日—26日", "六人环境、信息隔离与规则测试；五槽消息、开关、依据掩码与模板渲染的单元测试；随机和规则智能体可完整对局；确认课程题目"],
        ["沉默基线", "10月27日—11月2日", "S 的 PPO 学习曲线稳定并明显优于规则对手（关卡一）；M 配置通过玩具检查与通信诊断（关卡二）；F 配置 pilot 测速；冻结正式配置"],
        ["正式训练", "11月3日—16日", "核心层每个种子先训练 S 到分支点，再分出 M、F，各 3 种子；有余力启动逐项加入配置；checkpoint 与完整配置存档；HPC 有条件并行"],
        ["评估与机制", "11月17日—23日", "S → S、M → M、F → F、F → S 与测试时槽位消融；胜负分解、行为指标、RQ3 基线干预、RQ4 槽位分析"],
        ["分析与报告", "11月24日—30日", "区间与交叉对战；完成论文图表和演示；有余力再做自然语言扩展"],
    ], [1.13, 1.58, 4.0]
)
p("以上为六周实施安排，课程正式汇报时间与截止日期以教师确认的信息为准。第二阶段有两道关卡：关卡一是沉默基线本身，S 学不好时先排查环境、奖励、mask 与 credit assignment，不开始任何通信配置；关卡二是 M 配置的接收端影响，若接近零，先减少开启的槽位或缩小取值范围；不能直接把学习失败解释为“交流没有价值”。若到 11 月 9 日仍无法让通信被利用，则以沉默基线的完整分析（RQ1 胜负分解、RQ2 行为、RQ3 隐式信号）为主线，把通信学习失败作为一项结果报告，并结合文献讨论原因。")
p("最低交付包括：含两名平民、Morgana、Assassin、Merlin 和 Percival 的六人环境；带开关的五槽消息协议；随机／规则对手；沉默基线 S 与通信配置 M、F 各至少三种子的学习记录；S → S、M → M、F → F、F → S 与测试时槽位消融；RQ1 胜率及分解图；RQ2 行为比较图；RQ3 基线中的信息性分析及一项干预；RQ4 槽位使用、真实性与听者影响分析；实验复现说明、技术报告和课堂演示。五种子、逐项加入配置、从头训练的 F 对照、完整 cross-play 和自然语言扩展按资源追加。")
p("数据记录至少包含 game id、训练与评估 seed、通信配置及槽位开关、checkpoint、座位和角色映射、合法私有观察、提案、按槽位存储的消息元组及其模板渲染文本、公开投票、私有任务动作、公开任务结果、刺杀目标和最终胜负。私有真值日志与策略输入分开保存；同时保存配置、依赖版本、代码版本和分析脚本。")
p("可接受的失败与边界：若通信未被成功学会，报告优化失败及已验证的行为，不把它推广为游戏机制结论；若胜率差异很小，报告区间和可排除的效果范围；若 RQ3 只有相关性，保留为探索结果。有限预算下训练的 self-play 策略不等于 Nash equilibrium，本文测量的是特定学习系统的经验表现。")
p("最终成功标准是形成一个可复现、解释清楚的实验：说明哪些信息被允许、策略如何通过奖励学习、对照条件改变了什么、观察到的变化由哪些机制支持。结果不必证明沉默更偏向某一阵营，也不必强行发现隐式语言。")

h("参考文献")
ref("1", "Serrino, J., Kleiman-Weiner, M., Parkes, D. C., & Tenenbaum, J. B. (2019). Finding Friend and Foe in Multi-Agent Games. NeurIPS 2019.", "https://arxiv.org/abs/1906.02330")
ref("2", "Nakamura, Y. (2017). Is it Percival time yet? A preliminary analysis of Avalon gameplay and strategy. SIGBOVIK 2017. 人类游戏的探索性分析。", "https://sigbovik.org/2017/proceedings.pdf")
ref("3", "Light, J., Cai, M., Chen, W., Wang, G., Chen, X., Cheng, W., Yue, Y., & Hu, Z. (2025). Strategist: Self-improvement of LLM Decision Making via Bi-Level Tree Search. ICLR 2025. 六人对话与非对话设置见附录 O。", "https://openreview.net/forum?id=gfI9v7AbFg")
ref("4", "Light, J., Cai, M., Shen, S., & Hu, Z. (2023). AvalonBench: Evaluating LLMs Playing the Game of Avalon. NeurIPS 2023 Foundation Models for Decision Making Workshop.", "https://arxiv.org/abs/2310.05036")
ref("5", "Lan, Y., Hu, Z., Wang, L., Wang, Y., Ye, D., Zhao, P., Lim, E.-P., Xiong, H., & Wang, H. (2024). LLM-Based Agent Society Investigation: Collaboration and Confrontation in Avalon Gameplay. EMNLP 2024, 128–145.", "https://aclanthology.org/2024.emnlp-main.7/")
ref("6", "Lowe, R., Foerster, J., Boureau, Y.-L., Pineau, J., & Dauphin, Y. (2019). On the Pitfalls of Measuring Emergent Communication. AAMAS 2019.", "https://arxiv.org/abs/1903.05168")
ref("7", "Yu, C., Velu, A., Vinitsky, E., Gao, J., Wang, Y., Bayen, A., & Wu, Y. (2022). The Surprising Effectiveness of PPO in Cooperative Multi-Agent Games. NeurIPS 2022 Datasets and Benchmarks Track. 预印本始于 2021 年。", "https://arxiv.org/abs/2103.01955")
ref("8", "Schulman, J., Wolski, F., Dhariwal, P., Radford, A., & Klimov, O. (2017). Proximal Policy Optimization Algorithms. arXiv:1707.06347.", "https://arxiv.org/abs/1707.06347")
ref("9", "The Resistance: Avalon Rulebook. 游戏规则书，六人任务人数、表决、任务失败与刺杀规则；在线镜像。", "https://cdn.1j1ju.com/medias/a6/dc/c1-the-resistance-avalon-rulebook.pdf")
ref("10", "Light, J. et al. Avalon-LLM. AvalonBench 与 STRATEGIST 代码库。引擎位于 avalonbench_dev，README 说明目前只支持五人局；如参考其代码，记录具体 commit 并核对复用许可。", "https://github.com/jonathanmli/Avalon-LLM")
ref("11", "Brandizzi, N., Grossi, D., & Iocchi, L. (2021). RLupus: Cooperation through Emergent Communication in The Werewolf Social Deduction Game. Intelligenza Artificiale, 15(2), 55–70.", "https://arxiv.org/abs/2106.05018")
ref("12", "Meta Fundamental AI Research Diplomacy Team (FAIR), Bakhtin, A., Brown, N., et al. (2022). Human-level play in the game of Diplomacy by combining language models with strategic reasoning. Science, 378(6624), 1067–1074.", "https://www.science.org/doi/10.1126/science.ade9097")
ref("13", "Cao, K., Lazaridou, A., Lanctot, M., Leibo, J. Z., Tuyls, K., & Clark, S. (2018). Emergent Communication through Negotiation. ICLR 2018.", "https://arxiv.org/abs/1804.03980")
ref("14", "Jaques, N., Lazaridou, A., Hughes, E., Gulcehre, C., Ortega, P. A., Strouse, D. J., Leibo, J. Z., & de Freitas, N. (2019). Social Influence as Intrinsic Motivation for Multi-Agent Deep Reinforcement Learning. ICML 2019.", "https://arxiv.org/abs/1810.08647")
ref("15", "Rusu, A. A., Colmenarejo, S. G., Gulcehre, C., Desjardins, G., Kirkpatrick, J., Pascanu, R., Mnih, V., Kavukcuoglu, K., & Hadsell, R. (2016). Policy Distillation. ICLR 2016.", "https://arxiv.org/abs/1511.06295")
ref("16", "Foerster, J. N., Assael, Y. M., de Freitas, N., & Whiteson, S. (2016). Learning to Communicate with Deep Multi-Agent Reinforcement Learning. NeurIPS 2016.", "https://arxiv.org/abs/1605.06676")
p("[C1] CS7329 Fall 2026. Current Topics Presentation and Paper (Graduate Students). 本地保存的课程项目要求，2 页。")
p("[C2] Hahsler, M. Fall 2026 CS 7329 Reinforcement Learning 课程大纲；Introduction_to_Reinforcement_Learning 课程仓库本地备份。")
p("本地 references 文件夹保存核心 10 篇：[1]–[3]、[5]–[8]、[11]、[13]、[16]，逐篇要点与阅读顺序见其中的 README.md。阅读重点依次为 DeepRole 的信息结构与无语言设定、STRATEGIST 的规则附录与附录 O、PPO 与 MAPPO 的实现细节、RIAL 的消息训练方式、RLupus 的离散通信设计、Cao 等人关于 cheap talk 的结论，以及 Lowe 等人对通信测量的区分。")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run._r.addnext(fld)
footer.paragraph_format.space_after = Pt(0)
doc.save(OUT)
print(OUT)
print(f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}")
