from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "Silent_Avalon_Project_Brief.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.70), Inches(.70)
sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)
sec.footer_distance = Inches(.30)

def font_style(style, size, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + a), "Arial")
    fonts.set(qn("w:eastAsia"), "Heiti SC")
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + a), None)
    for el in list(style.element.iter(qn("w:pBdr"))):
        el.getparent().remove(el)

for name, size, bold in [
    ("Normal", 10.5, False), ("Title", 22, True), ("Subtitle", 12, False),
    ("Heading 1", 14, True), ("Heading 2", 11.5, True),
    ("List Bullet", 10.5, False), ("List Number", 10.5, False),
]:
    font_style(doc.styles[name], size, bold)

normal = doc.styles["Normal"].paragraph_format
normal.line_spacing = 1.18
normal.space_after = Pt(6)
normal.widow_control = True
for name in ("Heading 1", "Heading 2"):
    fmt = doc.styles[name].paragraph_format
    fmt.space_before = Pt(13 if name == "Heading 1" else 8)
    fmt.space_after = Pt(6)
    fmt.keep_with_next = True
doc.styles["Title"].paragraph_format.space_after = Pt(6)
doc.styles["Subtitle"].paragraph_format.space_after = Pt(12)

def p(text, boldlead=None, style=None):
    para = doc.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        para.add_run(boldlead).bold = True
        para.add_run(text[len(boldlead):])
    else:
        para.add_run(text)
    return para

def h(text, level=1):
    return doc.add_heading(text, level)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        borders.append(el)
    pr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, val in (("top", "85"), ("bottom", "85"), ("left", "105"), ("right", "105")):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:w"), val)
        e.set(qn("w:type"), "dxa")
        margins.append(e)
    pr.append(margins)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_idx, values in enumerate([headers] + rows):
        row = t.rows[0] if row_idx == 0 else t.add_row()
        cant = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant)
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.13
            run = para.add_run(value)
            run.font.size = Pt(9.5)
            if row_idx == 0:
                run.bold = True
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E6EDF3")
                cell._tc.get_or_add_tcPr().append(shade)
            if j == 0 and len(value) < 16:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.line_spacing = 1
    gap.add_run().font.size = Pt(2)
    return t

def ref(n, text, url):
    para = p(f"[{n}] {text}")
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.1
    for r in para.runs:
        r.font.size = Pt(9)
    para.add_run(" ")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "24567A")
    rpr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rpr.append(sz)
    r.append(rpr)
    tx = OxmlElement("w:t")
    tx.text = "原文链接"
    r.append(tx)
    link.append(r)
    para._p.append(link)

doc.core_properties.title = "六人阿瓦隆中显式交流的作用"
doc.core_properties.subject = "Silent Avalon Research Project Specification"
doc.core_properties.author = ""
doc.core_properties.keywords = "Avalon, MARL, Reinforcement Learning, Communication, CS7329"
doc.core_properties.comments = ""
p("六人阿瓦隆中显式交流的作用", style="Title")
p("Silent Avalon Research Project Specification", style="Subtitle")
p("CS7329 Reinforcement Learning　｜　项目规划　｜　2026年10月20日")

h("1 项目概述")
p("本项目研究六人隐藏身份游戏《阿瓦隆》中，移除显式交流如何改变阵营胜率、策略行为及信息传递。固定角色为两名平民、莫甘娜、刺客、梅林和派西维尔。研究采用同一套游戏规则与强化学习算法，分别训练可通信智能体与 Silent Avalon 智能体，并加入训练后禁言实验，区分渠道移除的即时影响与长期策略适应。")
p("建议的核心方案是小型 recurrent PPO 多智能体训练，加上可学习的离散消息信道；自然语言交流作为独立扩展实验。核心实验保留“显式交流对比沉默”的因果问题，能够利用本地 RTX 5070 Ti Laptop 完成开发和试验，并把学校 HPC 用于多随机种子训练。9B 模型不承担核心实验的端到端强化学习训练。", boldlead="建议的核心方案")
p("范围界定：离散消息是受限的显式通信，不能等同于开放式语言讨论。核心结论只适用于所定义的消息带宽、角色配置和训练所得策略。若最终论文要直接声称“自然语言被移除后的效果”，必须补充自然语言实验，并单独报告其结果。", boldlead="范围界定：")
h("研究问题", 2)
table(
    ["问题", "需要回答的内容", "主要证据"],
    [
        ["RQ1\nGame Balance", "禁言后，哪一阵营更容易获胜？", "好人最终胜率、三次任务成功概率、条件刺杀成功率"],
        ["RQ2\nStrategic Behavior", "组队、投票、破坏与身份隐藏如何调整？", "按角色和局面分层的行为指标及训练轨迹"],
        ["RQ3\nImplicit Signaling", "公开行为是否承担信息传递功能？", "信息增益、接收者干预、跨种子和跨搭档验证"],
    ], [1.12, 2.75, 2.84]
)
p("不预设沉默一定有利于坏人。交流可能帮助好人协调，也可能帮助坏人欺骗或识别梅林；任务阶段与刺杀阶段的效果还可能相互抵消。RQ1、RQ2 为核心研究目标，RQ3 为有明确证据标准的探索目标。即使没有显著胜率变化，或没有发现稳定信号，也可以形成有效结果。")

h("2 文献定位与课程关联")
p("DeepRole 研究五人阿瓦隆，明确不使用自然语言交流，以推断、搜索和学习的价值函数处理合作与对抗。[1] 因而“阿瓦隆可以在沉默条件下由智能体完成”已有先例；本项目应研究同一训练框架下通信条件变化带来的差异，而不是把 Silent Avalon 本身作为首次提出的环境。")
p("Is it Percival time yet 是 SIGBOVIK 2017 的初步人类游戏分析，报告 66 局采集、剔除 5 局后的 61 局数据。[2] 它可帮助提出刺杀、身份暴露和破坏协调的指标，但样本与角色配置存在限制，不足以直接推导禁言的因果效应。")
p("STRATEGIST 使用 LLM 生成和改进策略，并结合低层搜索进行决策，不能直接视为 PPO 式参数训练。[3] 尤其需要注意，附录 O 已描述六人局有对话与无对话实验。因此本项目的拟贡献是可复现的训练对照、策略适应分析及信号机制检验；单纯“六人局加禁言对照”不足以确立创新性。该文正文与附录的人类实验总局数表述不同，本文不把其汇总胜率当作可直接比较的基准。")
p("AvalonBench 及其代码库可作为环境实现的候选起点。[4,10] Lan 等人的六人 Avalon 研究采用包含 Percival 与 Morgana 的角色配置，并分析了协调、伪装和信息分享等行为，可用于借鉴行为分类。[5] 即使角色构成相同，具体规则、语言模型和实验规模仍须逐项核对，不应直接与本项目胜率拼成排行榜。")
p("拟贡献：提供一个在规则、算法和预算上受控的通信消融实验；把整体胜负分解为任务与刺杀两条路径；通过行为干预区分“动作包含信息”与“动作被用来影响其他智能体”。这是一项实验研究与机制分析项目，不以提出新的 RL 优化算法为必要条件。", boldlead="拟贡献：")
h("与课程内容的连接", 2)
p("课程中的 MDP 在这里扩展为 Partially Observable Stochastic Game。完整局面包括隐藏身份，单个玩家只能访问自己的 observation history。Actor 学习策略 π，critic 学习该玩家信息条件下的价值 V；终局胜负提供 return，TD 与 GAE 用于估计 advantage，PPO 用于更新策略。[7,8] Self-play 则使其他玩家的策略也随训练变化。")
p("课程项目要求动机与 RL 关联、相关研究、技术方法、文献或自有实验结果，以及局限性讨论。[C1] 本方案覆盖这些内容，并与仓库中的部分可观察性、函数近似、Actor–Critic 和 PPO 教学材料衔接。[C2]")

h("3 环境规则与信息边界")
h("核心六人配置", 2)
p("主实验固定为两名平民（Loyal Servants）、一名莫甘娜（Morgana）、一名刺客（Assassin）、一名梅林（Merlin）和一名派西维尔（Percival）。好人阵营为梅林、派西维尔及两名平民，坏人阵营为莫甘娜与刺客。可通信、沉默及训练后禁言条件均采用这一配置；派西维尔与莫甘娜的身份辨别及混淆机制属于核心研究范围。")
table(
    ["项目", "固定规范"],
    [
        ["任务与表决", "五轮任务人数为 2、3、4、3、4；至少四票赞成才通过，平票否决。同一任务连续五次提案被否决，坏人获胜。"],
        ["任务执行", "好人只能提交 Success；坏人可选 Success 或 Fail。六人局每轮一张 Fail 即失败。公开失败票总数，不公开个人任务票。"],
        ["终局判定", "三次任务失败即坏人胜；三次成功后由 Assassin 刺杀，命中 Merlin 则坏人胜，否则好人胜。"],
        ["身份信息", "所有人知道自己的角色。Merlin 知道两名坏人的座位，但不知道 Percival 的位置；Percival 看到 Merlin 与 Morgana 的两名候选者，但不知道谁是谁；Morgana 与 Assassin 互认坏人队友；两名平民没有额外身份信息。"],
        ["随机化", "每局随机分配角色和初始队长，随后队长按规则轮换；禁用 Lady of the Lake 等附加机制。"],
    ], [1.10, 5.61]
)
p("上述规则以游戏规则书为基准。[9] 座位编号不能固定对应角色；Percival 的两名候选者使用无身份标签的集合表示，不能用排列顺序泄露谁是 Merlin。投票与任务牌必须先分别提交、再一起揭示，防止执行顺序泄露未公开选择。仿真器完整状态仅用于裁判和离线分析。")
h("通信条件", 2)
p("Communicative Avalon 的核心信道是公开的离散消息：每次提案前、提案后，六名玩家按从队长开始的座位顺序各有一次发言机会；三次任务成功后，两名坏人各有一次公开的刺杀建议机会。消息可包括“玩家 j 是好人／坏人”“我声称某角色”“支持／反对当前提案”和“不发言”。不适用的类型按阶段屏蔽，刺杀建议允许指向候选玩家。所有合法消息均由策略选择，允许虚假声明，裁判不验证内容。")
p("Silent Avalon 使用相同游戏规则和观察格式，但所有显式消息均固定为 NULL，包括刺杀阶段。公开的队伍、已揭示投票、任务结果和失败票数仍然可见。两种条件都禁止私聊、跨玩家共享记忆及利用延迟、报错等额外信道。")
p("这种通信安排是有带宽限制的实验协议。记录消息词表、发言顺序和机会数并保持不变；不能把结果直接推广到任意长度、任意时间的自由讨论。给消息起了自然语言名称，也不代表智能体学到的含义一定与名称相同。")

h("4 强化学习方法")
h("策略与价值网络", 2)
p("采用 recurrent Independent PPO 作为首选基线：好人与坏人各使用一套小型 MLP＋GRU actor–critic，同阵营共享参数，输入包含自身角色和其合法私有信息。每个玩家独立维护 recurrent hidden state；共享网络权重不意味着共享其他玩家的记忆。先使用 128 维编码与 128 维 GRU 作为试验配置，并在正式比较前冻结架构。")
p("输入为公共游戏历史、自身历史行动、角色允许的信息及通信记录。输出使用分阶段 action heads：队长选择合法队伍，全员选择赞成或反对，坏人选择任务牌，Assassin 选择刺杀目标，通信阶段选择消息。采用 legal action mask 排除非法动作。六人局合法队伍最多 20 种，可直接枚举，无须生成任意文本动作。")
p("选择 recurrent policy 是为了利用历史信息，而非把当前观察错误地当成完整 Markov state。核心 critic 同样只接收该玩家可获得的信息；若日后采用 centralized critic，必须单独说明训练时的特权信息，且保证执行 actor 看不到它。")
h("奖励与更新", 2)
p("每个玩家在己方最终获胜时得到 +1、失败时得到 −1，中间奖励为 0；以有限回合未折扣终局收益为目标，取 γ＝1，避免新增通信步骤改变同一胜负的折扣权重。两阵营分别汇总损失，避免四名好人相对两名坏人造成隐含的训练权重差异。")
p("学习循环为：按当前策略采集完整对局 → 按玩家整理观察、合法动作、奖励和记忆序列 → 用价值估计与 GAE 计算 advantage → 用 PPO 更新行动头、消息头及 critic → 保存 checkpoint 并重复。通信动作也通过终局奖励学习。NULL 占位和强制动作不计入 actor 的可选动作损失；padding 不参与梯度。")
p("核心实验不增加“诚实发言”“成功传信”或“隐瞒身份”的额外奖励，避免把希望发现的策略直接写入奖励。是否出现信息交换和隐藏行为，由胜负激励下的学习结果决定。PPO、GAE 的具体超参数先在独立 pilot 中确定，之后两种条件共用；不能按测试胜率分别挑选最有利参数。")
h("Self play 与非平稳性", 2)
p("先训练两个阵营的当前策略，再维护历史 checkpoint 池，减少只适应最新对手的循环。建议 pilot 尝试每局以 50% 概率使用当前对手、25% 使用历史对手、25% 使用固定规则对手；正式实验锁定混合比例。对手的固定轨迹不参与当前策略的 on-policy 更新，同局队友仍保持独立观察。")
p("保留随机合法策略和固定规则策略作 sanity baselines。规则策略只利用合法信息，例如根据任务失败记录累计嫌疑、坏人以固定概率破坏、刺客按简单公开行为评分猜 Merlin。它们用于检查学习是否改善，不能代表人类或最优策略。IPPO 在此环境的有效性需要试验确认，不能由其在合作基准上的表现直接推断。[7]")

h("5 实验矩阵与统计计划")
table(
    ["实验", "训练条件", "评估条件", "解释"],
    [
        ["C → C", "可通信", "可通信", "主 baseline"],
        ["S → S", "沉默", "沉默", "适应沉默后的策略表现"],
        ["C → S", "可通信", "冻结权重后禁言", "立即失去通信渠道的影响"],
        ["S → S mask", "沉默", "削弱指定历史线索", "RQ3 的接收者依赖分析"],
    ], [1.0, 1.35, 1.75, 2.61]
)
p("主比较为 S → S 与 C → C。C → S 使用与 C → C 完全相同的 checkpoint 和评估初始条件，仅阻断消息产生和接收。它度量通信依赖，可能包含分布外输入效应，不等价于重新学会沉默游戏；与 S → S 的差异也不能简单解释为单一适应收益。")
p("训练预算：先完成每条件 2,000 局的功能与吞吐 pilot；依据耗时和学习曲线，在每次独立训练 20,000–100,000 局的候选范围内选择固定预算 B。该范围是规划参数，不是已验证的收敛需求。正式训练目标为每条件 5 个独立随机种子，即 10 次训练；资源不足时降为每条件 3 个，并将结果标为初步研究。")
p("两条件匹配对局数、网络容量、PPO 更新次数与超参数，报告有效决策数、消息动作数和实际训练耗时。额外消息动作意味着两条件的信息和动作预算本来就不同；匹配对局数不等于完全匹配计算量。因此同时给出性能对训练局数及实际时间的曲线。")
p("评估预算：各训练种子在 C → C、S → S、C → S 对应条件下分别冻结评估 1,000 局，目标共 15,000 局。配对使用相同角色分配和初始队长表，策略采样使用独立随机流；没有训练更新或跨局记忆。五个训练种子是主要重复单位，不能把 5,000 局当成 5,000 个独立训练模型。")
p("统计报告：先报告每个种子的指标，再汇总配对差值、效应量和 95% 区间。主结果使用按训练种子配对的分层 bootstrap，局内重复事件按完整游戏归组；种子数较少时明确区间不稳定。固定模型的单项胜率可另给 Wilson 区间，但不能用它替代训练不确定性。预先把好人最终胜率设为主指标，其余为机制与探索指标。")
p("同训练群体 self-play 只能说明该群体下的经验平衡。增加不同训练种子的好人和坏人交叉对战，以及与固定对手池的评估，以区分策略强弱、共同适应和通信效果。对阵双方都使用被测环境规则，且分别汇报每一阵营的表现。")

h("6 指标与可解释的结果")
h("RQ1 游戏平衡", 2)
p("核心指标为 ΔGood＝P(Good wins | S → S) − P(Good wins | C → C)。正值表示该实验设定下沉默更有利于好人，负值更有利于坏人。50% 只是阵营胜率对称的参考点，不意味着这款非对称游戏必须天然公平。")
p("为解释总胜率，分解为 P(Good wins)＝P(达到三次任务成功) × [1 − P(刺杀命中 Merlin | 达到三次任务成功)]。同时报告三次任务失败和连续五次否决导致的坏人获胜比例。刺杀成功率的分母只能是实际进入刺杀阶段的对局。")
p("任务成功率按任务序号报告，分母为实际执行该轮任务的次数，并报告到达该轮的局数；不能把未发生的后续任务记作失败。刺杀条件样本很少时报告宽区间，必要时按预先规定对所有条件等量增加评估。已知两名坏人位置时，在四名好人中均匀随机猜 Merlin 的命中率是 25%，可作刺杀参考基线。")
h("RQ2 策略行为", 2)
table(
    ["行为", "建议量化方式", "解释限制"],
    [
        ["组队与投票", "提议及通过队伍的好人比例；不同角色的赞成率；对含坏人队伍的赞成率", "按轮次、队伍人数、否决次数分层；最后机会的赞成不一定是判断错误"],
        ["任务破坏", "坏人在队时的 Fail 率；双坏人在队时零张、一张、两张 Fail 的分布", "放过任务可能用于伪装，两次破坏也不能自动判为协调失败"],
        ["梅林隐藏", "仅用公开历史训练身份 probe，评估识别 Merlin 的 log loss；结合任务与终局收益", "probe 难以识别不等于实际刺客一定被欺骗"],
        ["辨别与混淆", "Percival 对两名候选者的组队和投票偏好；Morgana 对这些偏好的影响", "控制局面与候选者历史；被支持不能直接等同于被认作 Merlin"],
        ["刺客推断", "条件刺杀准确率；对历史投票或组队记录的消融敏感度", "控制进入刺杀阶段的样本选择差异"],
    ], [1.0, 3.30, 2.41]
)
p("角色真值只在离线分析中用于计算这些指标。行为比较须覆盖早期、中期、最终 checkpoint；选取案例时使用明确抽样规则，同时展示支持与不支持假设的对局，不只挑选有趣的成功案例。")

h("7 隐式信号的证据标准")
p("公开行为与隐藏角色相关，只能证明行为包含可推断信息。要声称出现通信，还应证明接收者利用了这些行为；要进一步声称存在策略性信号，则需要发送者行为对接收者和收益的系统性作用。已有 emergent communication 研究说明，单看相关性或任务成绩容易产生误判。[6]")
p("第一层 信息性：使用独立测试对局训练和评估简单 probe。比较“公共局面＋队伍与任务结果”与“再加入个人投票序列”的隐藏身份预测效果。数据按完整对局和训练种子划分，避免同一对局片段进入训练与测试。预测提升只记为额外信息，不直接标为隐式通信。", boldlead="第一层 信息性：")
p("第二层 接收者使用：在固定 checkpoint 上回放同一接收者历史，将其他玩家的历史投票身份遮蔽，但保留总票数、提案结果与当前合法动作；从修改后的历史重新计算 GRU 状态，再比较接收者行动分布和后续收益。使用同长度占位作为 placebo，并以训练时见过遮蔽输入的辅助条件检验分布外影响。仅在推理时遮蔽导致性能下降，仍不足以单独证明通信。", boldlead="第二层 接收者使用：")
p("第三层 策略性信号：优先分析不改变表决通过结果的非关键投票。在可合法实现、保持票数与通过结果一致的反事实投票组合之间比较接收者反应，并检查信息持有者是否系统性选择能影响队友的模式。进一步改变搭档或其响应策略，观察发送者策略是否随之调整。若只完成前两层，应使用“行为线索被利用”的结论，避免宣称已经证明发送者意图或形成完整通信协议。", boldlead="第三层 策略性信号：")
p("重点分析 Percival 如何利用两名候选者的公开行为，以及 Morgana 是否形成稳定的混淆模式。稳健性检查包括角色与座位随机置换、未见过的角色分配和跨种子 cross-play。信号在不同种子间不兼容不等于群体内部没有信号。任务牌身份不公开，只能分析公开失败票数及其上下文。")
p("最低完成标准：至少完成信息性 probe 和一个可解释的接收者干预，清楚报告证据层级。若第三层证据不足，把 RQ3 作为探索性观察与后续研究方向，仍保留完整的 RQ1、RQ2 贡献。")

h("8 算力安排与自然语言扩展")
h("本地与 HPC 分工", 2)
p("本地 RTX 5070 Ti Laptop 优先运行规则引擎、并行环境、小型 recurrent PPO 和统计分析；先计时测量 CPU 仿真与 GPU 更新各自占比，不能默认瓶颈一定在 GPU。正式开跑前记录实际显存、RAM、功耗限制和软件版本，本规划不假设已验证该设备的吞吐或训练时长。")
p("建议为核心训练设置 40–60 GPU 小时的初始总预算上限，并用 pilot 修订 B：预计训练时间约为 2 × 种子数 × B × 每局平均训练耗时，另加调参与评估。若超预算，先降低单次训练量并保留重复种子；HPC 用于独立种子任务与扩展评估。核心项目不依赖持续取得 HPC 配额。")
p("复用 AvalonBench 的 engine 是候选路径，先核对六人规则、观察过滤、同时行动、五次否决和刺杀逻辑，再锁定代码版本。[10] 课程仓库提供算法学习与实现参考，但不是现成 Avalon 环境。此阶段不要求完整复现 DeepRole 或 STRATEGIST；若候选引擎改造过重，可按已固定规则实现轻量状态机。")
h("9B 模型与 API 的使用方式", 2)
p("自然语言扩展采用一个冻结的本地 9B 模型，或一个固定版本的低成本 API 模型，分别扮演六个独立玩家。每名玩家使用隔离的上下文和记忆；只共享模型权重。NL-C 允许同样时段的公开自由文本发言，NL-S 完全禁言，其他决策提示与预算保持一致。该实验验证语言代理中的方向是否类似，不能作为 MARL 策略已经学会自然语言的证据。")
p("先用 20 局 pilot 检查规则理解、格式合法率、上下文长度、单局 token 和时延，再视预算每条件完成 100–300 局。每次发言可先限制为 64 个输出 token，所有游戏动作通过结构化接口选择；对非法输出使用固定重试与回退规则并报告频率。不能因模型表现不佳而丢弃整局。")
p("9B 参数的原始权重在 FP16 下约需 18 GB，4-bit 下约需 4.5 GB，尚未包含量化元数据、KV cache 和运行开销。量化推理是否可行须实测；量化推理能运行不代表同一设备可承担六角色长轨迹的端到端 LLM RL 训练。优先单模型串行或受限批处理，避免常驻六份权重。")
p("API 成本按实际总输入与输出 token 分别乘以执行时价格计算；初始扩展预算可设为不超过 50 美元，超限则减少探索实验或转本地推理。本项目不需要先确定商业模型或付费方案。自由语言结果应单独展示，不能与离散通信实验混合平均。")
p("如果研究重点必须从第一天就是“自然语言＋可训练 RL”，可另选冻结 LLM 编码器加小型 RL 决策头的混合路线，但需额外设计消息理解、缓存与训练接口。这会改变时间和算力预算，建议作为后续项目，而非本课程最低交付条件。")

h("9 进度交付与失败处理")
table(
    ["阶段", "建议时间", "验收内容"],
    [
        ["规则与试运行", "10月20日—26日", "六人环境、信息隔离与规则测试；随机和规则智能体可完整对局；确认课程题目"],
        ["学习基线", "10月27日—11月2日", "Silent PPO 学习曲线；通信头接入；pilot 测速并冻结正式配置"],
        ["正式训练", "11月3日—16日", "C、S 多种子训练；checkpoint 与完整配置存档；HPC 有条件并行"],
        ["评估与机制", "11月17日—23日", "C → C、S → S、C → S；胜负分解、行为指标、RQ3 初步干预"],
        ["分析与报告", "11月24日—30日", "区间与交叉对战；完成论文图表和演示；有余力再做自然语言扩展"],
    ], [1.13, 1.58, 4.0]
)
p("以上为六周实施安排，课程正式汇报时间与截止日期以教师确认的信息为准。第二阶段是关键检查点：若策略仍接近随机或通信头没有被利用，应先排查信息、奖励、mask 与 credit assignment，再缩减模型和消息词表；不能直接把学习失败解释为“交流没有价值”。")
p("最低交付包括：含两名平民、Morgana、Assassin、Merlin 和 Percival 的六人环境；随机／规则基线及至少三种子 C、S 学习记录；三种主要评估条件；RQ1 胜率及分解图；RQ2 行为比较图；RQ3 信息性分析及一项干预；实验复现说明、技术报告和课堂演示。五种子、完整 cross-play 和自然语言扩展按资源追加。")
p("数据记录至少包含 game id、训练与评估 seed、checkpoint、座位和角色映射、合法私有观察、提案、消息、公开投票、私有任务动作、公开任务结果、刺杀目标和最终胜负。私有真值日志与策略输入分开保存；同时保存配置、依赖版本、代码版本和分析脚本。")
p("可接受的失败与边界：若通信未被成功学会，报告优化失败及已验证的行为，不把它推广为游戏机制结论；若胜率差异很小，报告区间和可排除的效果范围；若 RQ3 只有相关性，保留为探索结果。有限预算下训练的 self-play 策略不等于 Nash equilibrium，本文测量的是特定学习系统的经验表现。")
p("最终成功标准是形成一个可复现、解释清楚的实验：说明哪些信息被允许、策略如何通过奖励学习、对照条件改变了什么、观察到的变化由哪些机制支持。结果不必证明沉默更偏向某一阵营，也不必强行发现隐式语言。")

h("参考文献")
ref("1", "Serrino, J., Kleiman-Weiner, M., Parkes, D. C., & Tenenbaum, J. B. (2019). Finding Friend and Foe in Multi-Agent Games. NeurIPS 2019.", "https://arxiv.org/abs/1906.02330")
ref("2", "Nakamura, Y. (2017). Is it Percival time yet? A preliminary analysis of Avalon gameplay and strategy. SIGBOVIK 2017. 人类游戏的探索性分析。", "https://sigbovik.org/2017/proceedings.pdf")
ref("3", "Light, J., Cai, M., Chen, W., Wang, G., Chen, X., Cheng, W., Yue, Y., & Hu, Z. (2025). Strategist: Self-improvement of LLM Decision Making via Bi-Level Tree Search. ICLR 2025. 六人对话与非对话设置见附录 O。", "https://openreview.net/forum?id=gfI9v7AbFg")
ref("4", "Light, J., Cai, M., Shen, S., & Hu, Z. (2023). AvalonBench: Evaluating LLMs Playing the Game of Avalon. NeurIPS 2023 Foundation Models for Decision Making Workshop.", "https://arxiv.org/abs/2310.05036")
ref("5", "Lan, Y., Hu, Z., Wang, L., Wang, Y., Ye, D., Zhao, P., Lim, E.-P., Xiong, H., & Wang, H. (2024). LLM-Based Agent Society Investigation: Collaboration and Confrontation in Avalon Gameplay. EMNLP 2024, 128–145.", "https://aclanthology.org/2024.emnlp-main.7/")
ref("6", "Lowe, R., Foerster, J., Boureau, Y.-L., Pineau, J., & Dauphin, Y. (2019). On the Pitfalls of Measuring Emergent Communication. AAMAS 2019.", "https://arxiv.org/abs/1903.05168")
ref("7", "Yu, C., Velu, A., Vinitsky, E., Gao, J., Wang, Y., Bayen, A., & Wu, Y. (2022). The Surprising Effectiveness of PPO in Cooperative Multi-Agent Games. NeurIPS 2022 Datasets and Benchmarks Track. 预印本始于 2021 年。", "https://arxiv.org/abs/2103.01955")
ref("8", "Schulman, J., Wolski, F., Dhariwal, P., Radford, A., & Klimov, O. (2017). Proximal Policy Optimization Algorithms. arXiv:1707.06347.", "https://arxiv.org/abs/1707.06347")
ref("9", "The Resistance: Avalon Rulebook. 游戏规则书，六人任务人数、表决、任务失败与刺杀规则；在线镜像。", "https://cdn.1j1ju.com/medias/a6/dc/c1-the-resistance-avalon-rulebook.pdf")
ref("10", "Light, J. et al. Avalon-LLM. AvalonBench 与 STRATEGIST 代码库。候选环境入口为 avalonbench_dev；使用时记录具体 commit 并核对复用许可。", "https://github.com/jonathanmli/Avalon-LLM")
p("[C1] CS7329 Fall 2026. Current Topics Presentation and Paper (Graduate Students). 本地保存的课程项目要求，2 页。")
p("[C2] Hahsler, M. Fall 2026 CS 7329 Reinforcement Learning 课程大纲；Introduction_to_Reinforcement_Learning 课程仓库本地备份。")
p("文献阅读重点依次为 DeepRole 的信息结构与无语言设定、STRATEGIST 的方法及附录 O、通信测量中的相关性与行为影响区别，随后回到 PPO 和 recurrent actor–critic 实现。")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run._r.addnext(fld)
footer.paragraph_format.space_after = Pt(0)
doc.save(OUT)
print(OUT)
print(f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}")
