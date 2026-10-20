"""Shared python-docx layout for the Avalon research design and implementation spec."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

DOCS_DIR = Path(__file__).resolve().parents[1]  # generated .docx files go to docs/
TEXT_WIDTH = 6.71  # inches between margins


def _font_style(style, size, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + a), "Arial")
    fonts.set(qn("w:eastAsia"), "Heiti SC")
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + a), None)
    for el in list(style.element.iter(qn("w:pBdr"))):
        el.getparent().remove(el)


class DocBuilder:
    def __init__(self, title, subtitle, meta_line, subject, keywords):
        self.doc = doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
        sec.top_margin, sec.bottom_margin = Inches(.70), Inches(.70)
        sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)
        sec.footer_distance = Inches(.30)
        for name, size, bold in [
            ("Normal", 10.5, False), ("Title", 22, True), ("Subtitle", 12, False),
            ("Heading 1", 14, True), ("Heading 2", 11.5, True), ("List Bullet", 10.5, False),
        ]:
            _font_style(doc.styles[name], size, bold)
        normal = doc.styles["Normal"].paragraph_format
        normal.line_spacing = 1.18
        normal.space_after = Pt(6)
        normal.widow_control = True
        for name in ("Heading 1", "Heading 2"):
            fmt = doc.styles[name].paragraph_format
            fmt.space_before = Pt(13 if name == "Heading 1" else 8)
            fmt.space_after = Pt(6)
            fmt.keep_with_next = True
        doc.styles["List Bullet"].paragraph_format.space_after = Pt(3)
        doc.styles["Title"].paragraph_format.space_after = Pt(6)
        doc.styles["Subtitle"].paragraph_format.space_after = Pt(12)

        doc.core_properties.title = title
        doc.core_properties.subject = subject
        doc.core_properties.author = ""
        doc.core_properties.keywords = keywords
        doc.core_properties.comments = ""
        self.p(title, style="Title")
        self.p(subtitle, style="Subtitle")
        self.p(meta_line)

        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run()
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run._r.addnext(fld)
        footer.paragraph_format.space_after = Pt(0)

    def p(self, text, boldlead=None, style=None):
        para = self.doc.add_paragraph(style=style)
        if boldlead and text.startswith(boldlead):
            para.add_run(boldlead).bold = True
            para.add_run(text[len(boldlead):])
        else:
            para.add_run(text)
        return para

    def h(self, text, level=1):
        return self.doc.add_heading(text, level)

    def bullets(self, items):
        """items: str, or (boldlead, rest) tuples rendered as a bold lead-in."""
        for item in items:
            if isinstance(item, tuple):
                lead, rest = item
                self.p(lead + rest, boldlead=lead, style="List Bullet")
            else:
                self.p(item, style="List Bullet")

    def table(self, headers, rows, widths, keep=True, center_first=True):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
        pr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + edge)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "D9D9D9")
            borders.append(el)
        pr.append(borders)
        margins = OxmlElement("w:tblCellMar")
        for edge, val in (("top", "85"), ("bottom", "85"), ("left", "105"), ("right", "105")):
            e = OxmlElement("w:" + edge)
            e.set(qn("w:w"), val)
            e.set(qn("w:type"), "dxa")
            margins.append(e)
        pr.append(margins)
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for row_idx, values in enumerate([headers] + rows):
            row = t.rows[0] if row_idx == 0 else t.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for j, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = Inches(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.13
                if keep and row_idx < len(rows):
                    para.paragraph_format.keep_with_next = True
                run = para.add_run(value)
                run.font.size = Pt(9.5)
                if row_idx == 0:
                    run.bold = True
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E6EDF3")
                    cell._tc.get_or_add_tcPr().append(shade)
                if center_first and j == 0 and len(value) < 16:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        gap = self.doc.add_paragraph()
        gap.paragraph_format.space_after = Pt(0)
        gap.paragraph_format.space_before = Pt(0)
        gap.paragraph_format.line_spacing = 1
        gap.add_run().font.size = Pt(2)
        return t

    def ref(self, n, text, url=None):
        para = self.p(f"[{n}] {text}")
        para.paragraph_format.space_after = Pt(5)
        para.paragraph_format.line_spacing = 1.1
        for r in para.runs:
            r.font.size = Pt(9)
        if url:
            para.add_run(" ").font.size = Pt(9)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "24567A")
            rpr.append(color)
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), "18")
            rpr.append(sz)
            r.append(rpr)
            tx = OxmlElement("w:t")
            tx.text = "原文链接"
            r.append(tx)
            link.append(r)
            para._p.append(link)

    def save(self, name):
        out = DOCS_DIR / name
        self.doc.save(out)
        print(out)
        print(f"Paragraphs: {len(self.doc.paragraphs)}; tables: {len(self.doc.tables)}")
        return out
